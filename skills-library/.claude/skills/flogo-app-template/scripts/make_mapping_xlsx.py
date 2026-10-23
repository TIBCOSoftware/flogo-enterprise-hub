"""Create the sample mapping workbooks used as template inputs.

Each workbook has two sheets:
  Event     - how the canonical Kafka event is built from the REST request
  Response  - what the API replies (HTTP 202 body)

Columns: Target Field | Type | Mapping Rule | Description
Mapping-rule notation (translated deterministically by generate_app.py):
  request.<field>      field of the incoming REST body
  event.<field>        field of the event built on the Event sheet (Response sheet only)
  [Property.Name]      Flogo app property
  "text"               fixed literal
  UPPER() LOWER() TRIM() concat() newUUID() now() asString() asInt() asNumber()
  =<flogo expression>  raw Flogo expression (escape hatch)

Usage:  python scripts/make_mapping_xlsx.py [--out <inputs-dir>]   (default ./inputs; writes <dir>/flight|baggage/*-mapping.xlsx)
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INPUTS = os.path.join(os.getcwd(), "inputs")   # never inside the skill folder

HEADER = ["Target Field", "Type", "Mapping Rule", "Description"]

FLIGHT_EVENT = [
    ("eventId", "string", "newUUID()", "Unique id for every published event"),
    ("eventType", "string", '"FlightStatusChanged"', "Fixed event type"),
    ("eventTime", "string", "now()", "Time the gateway produced the event"),
    ("schemaVersion", "string", '"1.0"', "Canonical event schema version"),
    ("flightKey", "string", 'concat(UPPER(request.carrier), request.flightNumber, "-", request.departureDate)',
     "Business key, also used as the Kafka message key"),
    ("flight.carrier", "string", "UPPER(request.carrier)", "IATA carrier code, upper case"),
    ("flight.number", "string", "request.flightNumber", ""),
    ("flight.departureDate", "string", "request.departureDate", "Local departure date (yyyy-mm-dd)"),
    ("route.origin", "string", "UPPER(request.origin)", "IATA airport code"),
    ("route.destination", "string", "UPPER(request.destination)", "IATA airport code"),
    ("status.code", "string", "UPPER(request.status)", "SCHEDULED / DELAYED / BOARDING / DEPARTED / CANCELLED"),
    ("status.gate", "string", "request.gate", ""),
    ("status.delayMinutes", "integer", "request.delayMinutes", ""),
    ("status.delayReason", "string", "TRIM(request.delayReason)", ""),
    ("source", "string", "request.updatedBy", "Operations system or desk that sent the update"),
]
FLIGHT_RESPONSE = [
    ("eventId", "string", "event.eventId", "Id of the published event"),
    ("status", "string", '"ACCEPTED"', ""),
    ("topic", "string", "[Event.Topic]", "Topic the event was published to"),
    ("flightKey", "string", "event.flightKey", ""),
]

BAGGAGE_EVENT = [
    ("eventId", "string", "newUUID()", "Unique id for every published event"),
    ("eventType", "string", '"BaggageScanned"', "Fixed event type"),
    ("eventTime", "string", "now()", "Time the gateway produced the event"),
    ("schemaVersion", "string", '"1.0"', "Canonical event schema version"),
    ("bagKey", "string", 'concat(request.bagTag, "-", request.departureDate)',
     "Business key, also used as the Kafka message key"),
    ("bag.tag", "string", "request.bagTag", "10-digit licence plate"),
    ("bag.weightKg", "number", "request.weightKg", ""),
    ("bag.bookingRef", "string", "UPPER(request.bookingRef)", "PNR record locator"),
    ("flight.carrier", "string", "UPPER(request.carrier)", ""),
    ("flight.number", "string", "request.flightNumber", ""),
    ("flight.departureDate", "string", "request.departureDate", ""),
    ("scan.station", "string", "UPPER(request.station)", "IATA airport code"),
    ("scan.point", "string", "UPPER(request.scanPoint)", "CHECKIN / SORTER / LOADED / ARRIVED"),
    ("scan.time", "string", "request.scanTime", ""),
]
BAGGAGE_RESPONSE = [
    ("eventId", "string", "event.eventId", ""),
    ("status", "string", '"ACCEPTED"', ""),
    ("topic", "string", "[Event.Topic]", ""),
    ("bagKey", "string", "event.bagKey", ""),
]


def write_sheet(ws, rows):
    ws.append(HEADER)
    for r in rows:
        ws.append(list(r))
    fill = PatternFill("solid", fgColor="1F3864")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    for col, width in zip("ABCD", (24, 10, 78, 52)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADER))}{ws.max_row}"


def make(path, event_rows, response_rows):
    wb = Workbook()
    write_sheet(wb.active, event_rows)
    wb.active.title = "Event"
    write_sheet(wb.create_sheet("Response"), response_rows)
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else INPUTS
    for sub in ("flight", "baggage"):
        os.makedirs(os.path.join(out, sub), exist_ok=True)
    make(os.path.join(out, "flight", "flight-mapping.xlsx"), FLIGHT_EVENT, FLIGHT_RESPONSE)
    make(os.path.join(out, "baggage", "baggage-mapping.xlsx"), BAGGAGE_EVENT, BAGGAGE_RESPONSE)
    sys.exit(0)
