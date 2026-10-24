"""Shared helpers for the Flogo app-template demo (harvest.py / generate_app.py).

Everything here is deterministic: the same inputs always produce the same
parameters, mapping rows and fda command lines.
"""
import glob
import json
import os
import re
import shlex
import subprocess
import sys

HERE = os.path.dirname(os.path.abspath(__file__))


def _find_config():
    """Walk up from this script and from the cwd to the skills-library `.claude/skills/config.md`."""
    for start in (HERE, os.getcwd()):
        d = start
        while True:
            for cand in (os.path.join(d, "config.md"), os.path.join(d, ".claude", "skills", "config.md")):
                if os.path.basename(os.path.dirname(cand)) == "skills" and os.path.exists(cand):
                    return cand
            parent = os.path.dirname(d)
            if parent == d:
                break
            d = parent
    return ""


CONFIG_MD = _find_config()

ANSI = re.compile(r"\x1b\[[0-9;]*m")

# Activity / connection names fixed by the template pattern (see pattern-guide.md).
MAPPER = "BuildEvent"
PROPERTIES = {"Event.Topic", "Kafka.EventBus.Brokers", "HTTP.Port"}


class TemplateError(Exception):
    pass


# --------------------------------------------------------------------------- config / CLI

def read_config():
    """Return {KEY: value} from the `| KEY | `value` |` rows of config.md (empty if missing)."""
    cfg = {}
    if CONFIG_MD:
        for line in open(CONFIG_MD, encoding="utf-8"):
            m = re.match(r"\|\s*([A-Z_]+)\s*\|\s*`([^`]*)`", line)
            if m:
                cfg[m.group(1)] = m.group(2)
    return cfg


def resolve_fda():
    """FDA env var -> newest tibco.flogo-* VS Code extension -> config.md FLOGODESIGN_CLI_PATH."""
    if os.environ.get("FDA"):
        return os.environ["FDA"]
    exe = "flogodesign-cli.exe" if os.name == "nt" else "flogodesign-cli"
    home = os.path.expanduser("~")
    for root in (".vscode", ".vscode-insiders", ".vscode-server"):
        hits = glob.glob(os.path.join(home, root, "extensions", "tibco.flogo-*", "bin", exe))
        if hits:
            return max(hits, key=os.path.getmtime)
    p = (read_config().get("FLOGODESIGN_CLI_PATH") or "").strip('"')   # full exe path or its bin directory
    if p:
        for cand in (p, os.path.join(p, exe)):
            if os.path.isfile(cand):
                return cand
    raise TemplateError("flogodesign-cli (fda) not found - install the TIBCO Flogo VS Code extension")


def fda_env():
    env = dict(os.environ)
    env["MSYS_NO_PATHCONV"] = "1"
    env["FLOGO_DESIGN_ENABLE_HISTORY"] = "true"
    return env


def fda_version(fda):
    out = subprocess.run([fda, "version"], capture_output=True, text=True, encoding="utf-8", errors="replace")
    return ANSI.sub("", out.stdout + out.stderr).strip()


def run_fda(fda, argv, flogo, cwd):
    """Run one fda command; return (ok, last_message)."""
    p = subprocess.run([fda, *argv, "-f", flogo], cwd=cwd, env=fda_env(), capture_output=True,
                       text=True, encoding="utf-8", errors="replace")
    out = ANSI.sub("", p.stdout + p.stderr)
    lines = [l for l in out.splitlines()
             if l.strip() and "DEP0169" not in l and "trace-deprecation" not in l and "History Enabled" not in l]
    # `fda cm` exits 0 and prints "(INFO) N error(s)" when mappings fail, so check its summary as well.
    ok = p.returncode == 0 and "(ERROR)" not in out and not re.search(r"\b[1-9]\d* error\(s\)", out)
    return ok, (lines[-1] if lines else "").replace("[Flogo Design] ", "").strip()


# --------------------------------------------------------------------------- OpenAPI

def _lookup(spec, ref):
    """Follow a local JSON pointer such as #/components/schemas/Order."""
    if not ref.startswith("#/"):
        raise TemplateError(f"only local $refs are supported, not {ref}")
    node = spec
    for part in ref[2:].split("/"):
        part = part.replace("~1", "/").replace("~0", "~")
        if not isinstance(node, dict) or part not in node:
            raise TemplateError(f"$ref {ref} does not resolve")
        node = node[part]
    return node


def _resolve(node, spec, seen=(), is_properties=False):
    """Inline every $ref. A recursive $ref stops as an empty (free-form) schema. Examples are dropped, but a field
    that happens to be NAMED `example` (a key of a `properties` map) is kept."""
    if isinstance(node, dict):
        if "$ref" in node:
            ref = node["$ref"]
            return {} if ref in seen else _resolve(_lookup(spec, ref), spec, seen + (ref,))
        return {k: _resolve(v, spec, seen, (not is_properties) and k == "properties")
                for k, v in node.items() if is_properties or k != "example"}
    if isinstance(node, list):
        return [_resolve(v, spec, seen) for v in node]
    return node


def _json_schema(holder, spec):
    """The application/json schema of a request body or response object (either may be a $ref)."""
    if isinstance(holder, dict) and "$ref" in holder:
        holder = _lookup(spec, holder["$ref"])
    for media, content in ((holder or {}).get("content") or {}).items():
        if media.split(";")[0].strip().lower() == "application/json" and isinstance(content, dict):
            return content.get("schema")
    return None


def read_spec(path, operation_id):
    """Return the facts the template needs from an OpenAPI 3 (JSON) spec."""
    spec = json.load(open(path, encoding="utf-8"))
    for api_path, ops in spec.get("paths", {}).items():
        for method, op in ops.items():
            if isinstance(op, dict) and op.get("operationId") == operation_id:
                body = _json_schema(op.get("requestBody"), spec)
                if body is None:
                    raise TemplateError(f"operation {operation_id} has no application/json request body")
                replies = {c: s for c, s in ((c, _json_schema(r, spec)) for c, r in op.get("responses", {}).items())
                           if str(c).startswith("2") and s}
                if not replies:
                    raise TemplateError(f"operation {operation_id} has no 2xx application/json response")
                code = "202" if "202" in replies else sorted(replies)[0]   # the pattern replies 202 Accepted
                return {
                    "title": spec["info"]["title"],
                    "path": api_path,
                    "method": method.upper(),
                    "request": _resolve(body, spec),
                    "replyCode": code,
                    "reply": _resolve(replies[code], spec),
                }
    raise TemplateError(f"operationId '{operation_id}' not found in {path}")


def trigger_name(title):
    """Name `fda cas` gives the tr_rest trigger: the spec title with spaces as `_` (fda 0.9.3 keeps every other
    character). fda cannot address a trigger whose name contains a `.`, so such titles are rejected."""
    if "." in title:
        raise TemplateError(f"OpenAPI info.title '{title}' contains '.', which fda cannot use in a trigger name")
    return title.replace(" ", "_")


# --------------------------------------------------------------------------- mapping workbook

FUNCS = {
    "UPPER": "string.toUpper", "LOWER": "string.toLower", "TRIM": "string.trim",
    "concat": "string.concat", "CONCAT": "string.concat", "LENGTH": "string.length",
    "SUBSTRING": "string.substring", "REPLACE": "string.replaceAll",
    "newUUID": "utils.uuid", "now": "datetime.now",
    "asString": "coerce.toString", "asInt": "coerce.toInt", "asNumber": "coerce.toFloat64",
}
QUALIFIED_PKGS = {"string", "datetime", "coerce", "utils", "number", "array", "json", "math", "boolean"}
JSON_TYPES = {"string", "integer", "number", "boolean"}


def read_workbook(path):
    from openpyxl import load_workbook  # only needed when a workbook is read

    # Formula text, not cached results: a "=<raw Flogo expression>" rule is stored by Excel as a formula.
    wb = load_workbook(path, data_only=False)
    sheets = {}
    for name in ("Event", "Response"):
        if name not in wb.sheetnames:
            raise TemplateError(f"{os.path.basename(path)}: missing sheet '{name}'")
        rows = []
        for i, r in enumerate(wb[name].iter_rows(min_row=2, values_only=True), start=2):
            target, typ, rule = (r + (None, None, None))[:3]
            if not target:
                continue
            rows.append({"row": i, "target": str(target).strip(), "type": str(typ or "string").strip().lower(),
                         "rule": "" if rule is None else str(rule).strip()})
        sheets[name] = rows
    return sheets


def check_path(schema, dotted, what, where):
    """Fail unless every segment of a dotted field path is declared in the (resolved) JSON schema.
    Stops at a node that declares no properties (free-form object, array, allOf/oneOf): nothing to check against."""
    node, walked = schema or {}, []
    for seg in dotted.split("."):
        props = node.get("properties") if isinstance(node, dict) else None
        if props is None:
            return
        if seg not in props:
            at = ".".join(walked) or "the top level"
            raise TemplateError(f"{where}: '{dotted}' is not in the {what} ('{seg}' is not a field of {at}; "
                                f"fields there: {', '.join(sorted(props))})")
        node = props[seg]
        walked.append(seg)


def translate(rule, typ, where, request_schema, event_fields):
    """Translate one friendly mapping rule into (value, extra_fda_args)."""
    if not rule or rule == "=":
        raise TemplateError(f"{where}: the Mapping Rule is empty")
    if rule.startswith("="):
        return rule, []
    m = re.fullmatch(r'"([^"]*)"', rule)
    if m:
        return m.group(1), []
    if typ in ("integer", "number") and re.fullmatch(r"-?\d+(\.\d+)?", rule):
        return rule, ["--type", "number"]
    if typ == "boolean" and rule.lower() in ("true", "false"):
        return rule.lower(), ["--type", "boolean"]

    out, i = [], 0
    while i < len(rule):
        ch = rule[i]
        if ch == '"':
            j = i + 1
            while j < len(rule) and rule[j] != '"':
                j += 2 if rule[j] == "\\" else 1   # a backslash escapes the next character (\" and \\)
            if j >= len(rule):
                raise TemplateError(f"{where}: unterminated string literal in {rule!r}")
            out.append(rule[i:j + 1])
            i = j + 1
        elif ch == "[":
            j = rule.find("]", i)
            if j < 0:
                raise TemplateError(f"{where}: '[' without a closing ']' in {rule!r}")
            prop = rule[i + 1:j]
            if prop not in PROPERTIES:
                raise TemplateError(f"{where}: unknown app property [{prop}] (known: {', '.join(sorted(PROPERTIES))})")
            out.append(f'$property["{prop}"]')
            i = j + 1
        elif ch.isalpha() or ch == "_":
            m = re.match(r"[A-Za-z_][\w.]*", rule[i:])
            ident = m.group(0)
            i += len(ident)
            rest = rule[i:].lstrip()
            if rest.startswith("("):
                if ident in FUNCS:
                    out.append(FUNCS[ident])
                elif ident.split(".")[0] in QUALIFIED_PKGS and "." in ident:
                    out.append(ident)
                else:
                    raise TemplateError(f"{where}: unknown function {ident}() (known: {', '.join(sorted(FUNCS))})")
            elif ident.startswith("request."):
                check_path(request_schema, ident[len("request."):], "API request body", where)
                out.append("$flow.body." + ident[len("request."):])
            elif ident.startswith("event."):
                if event_fields is None:
                    raise TemplateError(f"{where}: event.* can only be used on the Response sheet")
                if ident[len("event."):] not in event_fields:
                    raise TemplateError(f"{where}: {ident} is not a target on the Event sheet")
                out.append(f"$activity[{MAPPER}].output." + ident[len("event."):])
            elif ident in ("true", "false", "nil"):
                out.append(ident)
            else:
                raise TemplateError(f"{where}: don't know what '{ident}' refers to "
                                    f"(use request.<field>, event.<field>, [Property] or \"literal\")")
        else:
            out.append(ch)
            i += 1
    return "=" + "".join(out), []


def event_schema(rows):
    """Nested JSON schema built from the dotted Event-sheet targets, in sheet order."""
    root = {"$schema": "http://json-schema.org/draft-04/schema#", "type": "object", "properties": {}}
    seen = {}   # dotted target -> row, for duplicate / leaf-vs-object conflicts
    for r in rows:
        where = f"Event row {r['row']} ({r['target']})"
        if r["type"] not in JSON_TYPES:
            raise TemplateError(f"{where}: type '{r['type']}' must be one of {sorted(JSON_TYPES)}")
        parts = r["target"].split(".")
        if not all(re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", p) for p in parts):
            raise TemplateError(f"{where}: each part of a target must be a field name (letters, digits, _)")
        if r["target"] in seen:
            raise TemplateError(f"{where}: duplicate target (also on row {seen[r['target']]})")
        for k in range(1, len(parts)):
            if ".".join(parts[:k]) in seen:
                raise TemplateError(f"{where}: '{'.'.join(parts[:k])}' is a field on row "
                                    f"{seen['.'.join(parts[:k])]}, so it cannot also be an object")
        clash = next((t for t in seen if t.startswith(r["target"] + ".")), None)
        if clash:
            raise TemplateError(f"{where}: '{r['target']}' is an object on row {seen[clash]} ({clash}), "
                                "so it cannot also be a field")
        seen[r["target"]] = r["row"]
        node = root
        for p in parts[:-1]:
            node = node["properties"].setdefault(p, {"type": "object", "properties": {}})
        node["properties"][parts[-1]] = {"type": r["type"]}
    return root


# --------------------------------------------------------------------------- parameters

def derive(params_path):
    """All template parameters and @foreach rows, derived from one <app>.params.json."""
    base = os.path.dirname(os.path.abspath(params_path))
    p = json.load(open(params_path, encoding="utf-8"))
    for k in ("appName", "description", "swagger", "operationId", "mappingWorkbook",
              "eventSchemaName", "keyField", "topic", "brokers", "httpPort"):
        if k not in p:
            raise TemplateError(f"{params_path}: missing '{k}'")
    if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", p["appName"]):
        raise TemplateError("appName must be letters/digits/underscore and start with a letter")
    if not re.fullmatch(r"\d{1,5}", str(p["httpPort"])) or not 1 <= int(p["httpPort"]) <= 65535:
        raise TemplateError(f"httpPort must be a whole number from 1 to 65535 (got {p['httpPort']!r})")

    spec_path = os.path.join(base, p["swagger"])
    spec = read_spec(spec_path, p["operationId"])
    sheets = read_workbook(os.path.join(base, p["mappingWorkbook"]))
    event_rows, event_targets = [], set()
    for r in sheets["Event"]:
        v, extra = translate(r["rule"], r["type"], f"Event row {r['row']} ({r['target']})", spec["request"], None)
        event_rows.append({"target": r["target"], "value": v, "extra": extra})
        event_targets.add(r["target"])
    if p["keyField"] not in event_targets:
        raise TemplateError(f"keyField '{p['keyField']}' is not a target on the Event sheet")

    reply_rows = []
    for r in sheets["Response"]:
        check_path(spec["reply"], r["target"], "API 2xx response", f"Response row {r['row']}")
        v, extra = translate(r["rule"], r["type"], f"Response row {r['row']} ({r['target']})", spec["request"],
                             event_targets)
        reply_rows.append({"target": r["target"], "value": v, "extra": extra})

    compact = lambda o: json.dumps(o, separators=(",", ":"))
    values = {
        "APP_NAME": p["appName"],
        "APP_DESCRIPTION": p["description"],
        "SPEC_NAME": p["appName"] + "Api",
        "SPEC_FILE": os.path.basename(spec_path),
        "TRIGGER": trigger_name(spec["title"]),
        "FLOW": f"{p['operationId'].replace(' ', '_')}_{spec['method']}",   # as `fda cas` names the flow
        "HTTP_PORT": str(int(p["httpPort"])),
        "BROKERS": p["brokers"],
        "TOPIC": p["topic"],
        "EVENT_SCHEMA_NAME": p["eventSchemaName"],
        "EVENT_SCHEMA": compact(event_schema(sheets["Event"])),
        "KEY_FIELD": p["keyField"],
        "REPLY_SCHEMA": compact({"type": "object", "properties": {"body": spec["reply"]}}),
    }
    return {"values": values, "rows": {"EVENT_MAPPING": event_rows, "REPLY_MAPPING": reply_rows},
            "spec_path": spec_path, "spec": spec}


# --------------------------------------------------------------------------- history parsing

# Positional arity per fda task; the LAST positional is greedy (history joins argv with spaces).
ARITY = {"cp": 2, "csp": 2, "cas": 2, "cap": 3, "rap": 1, "cc": 2, "sa": 3, "ca": 4, "aa": 4, "cl": 3,
         "rl": 3, "sms": 2, "mm": 2, "rmm": 1, "cs": 2, "rs": 1, "cf": 2, "rf": 1, "ct": 3, "cth": 3,
         "wth": 2, "ai": 1, "ri": 1, "cat": 3, "us": 2, "ums": 2}
# Flags and whether they take a value, from `fda help <task>` (fda 0.9.3). A flag missing here is glued to its value
# when a history entry is re-split, so keep this complete when fda adds options.
FLAGS = {"-f": 1, "--file": 1, "-H": 0, "--history": 0, "-j": 0, "--json": 0, "--force": 0, "--doNotLink": 0,
         "--specFile": 1, "-C": 1, "--connection": 1, "--configType": 1, "--json-schema": 1, "--json-schema-file": 1,
         "--jsonSchemaFile": 1, "--json-value-to-schema": 1, "--json-value-to-schema-file": 1, "--type": 1, "-ty": 1,
         "--jsonValue": 1, "-jv": 1, "--jsonFile": 1, "--jsonValueFile": 1, "-jvf": 1, "--foreach": 1, "-fe": 1,
         "--as": 1, "--input": 1, "--output": 1, "--inputs-only": 0, "--reply-only": 0, "--restResourcePath": 1,
         "--restHandlerMethod": 1, "--mcpHandlerType": 1, "--mcpHandlerName": 1, "--mcpHandlerDescription": 1,
         "--field": 1, "--alias": 1}


def parse_history_command(cmd):
    """Re-split a history entry (argv joined by spaces, quotes lost) into argv."""
    toks = ANSI.sub("", cmd).split(" ")
    task, toks = toks[0], toks[1:]
    n = ARITY.get(task)
    if n is None:
        raise TemplateError(f"harvest: no arity known for task '{task}' in: {cmd}")
    argv, i, pos = [task], 0, 0
    while i < len(toks):
        t = toks[i]
        if t in FLAGS:
            k = FLAGS[t]
            if k == 0:
                argv.append(t)
                i += 1
            else:  # flag value is greedy up to the next known flag
                j = i + 1
                while j < len(toks) and toks[j] not in FLAGS:
                    j += 1
                argv += [t, " ".join(toks[i + 1:j])]
                i = j
        elif pos < n - 1:
            argv.append(t)
            pos += 1
            i += 1
        else:  # last positional: greedy
            j = i
            while j < len(toks) and toks[j] not in FLAGS:
                j += 1
            argv.append(" ".join(toks[i:j]))
            pos += 1
            i = j
    return argv


def strip_file_flags(argv):
    out, i = [], 0
    while i < len(argv):
        if argv[i] in ("-f", "--file"):
            i += 2
        elif argv[i] in ("-H", "--history"):
            i += 1
        else:
            out.append(argv[i])
            i += 1
    return out


def shell_line(argv):
    return " ".join(shlex.quote(a) for a in argv)


def rel(path):
    """Path relative to the cwd for display; absolute when that is impossible (another drive on Windows)."""
    try:
        return os.path.relpath(path)
    except ValueError:
        return os.path.abspath(path)


def utf8_console():
    """Print non-ASCII values (descriptions, literals) instead of crashing on a cp1252 pipe on Windows."""
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")


def die(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)
